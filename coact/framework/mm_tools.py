# Copyright 2025 The OPPO Inc. Personal AI team. All rights reserved.
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

import base64
import json
import logging
import mimetypes
import os
import time
import uuid

import requests
from openpyxl import load_workbook
from PIL import Image
from pptx import Presentation

from .mm_tools_utils import MarkdownConverter
from .models import MessageRole, Model
from .tools import Tool

logger = logging.getLogger(__name__)


class VisualInspectorTool(Tool):
    name = "inspect_file_as_image"
    description = """
You cannot load files directly: use this tool to process image files and answer related questions.
This tool supports the following image formats: [".jpg", ".jpeg", ".png", ".gif", ".bmp"]. For other file types, use the appropriate inspection tool."""

    inputs = {
        "file_path": {
            "description": "The path to the file you want to read as an image. Must be a '.something' file, like '.jpg','.png','.gif'. If it is text, use the text_inspector tool instead! If it is audio, use the audio_inspector tool instead! DO NOT use this tool for an HTML webpage: use the web_search tool instead!",
            "type": "string",
        },
        "question": {
            "description": "[Optional]: Your question about the image content. Provide as much context as possible. Do not pass this parameter if you just want to get a description of the image.",
            "type": "string",
            "nullable": True,
        },
    }
    output_type = "string"

    def __init__(self, model: Model, text_limit: int):
        super().__init__()
        self.model = model
        self.text_limit = text_limit
        self.gpt_key = os.getenv("OPENAI_API_KEY")
        self.gpt_url = os.getenv("OPENAI_URL")

    def _validate_file_type(self, file_path: str):
        if not any(
            file_path.lower().endswith(ext) for ext in [".jpg", ".jpeg", ".png", ".gif", ".bmp"]
        ):
            raise ValueError(
                "Unsupported file type. Use the appropriate tool for text/audio files."
            )

    def _resize_image(self, image_path: str) -> str:
        img = Image.open(image_path)
        width, height = img.size
        img = img.resize((int(width / 2), int(height / 2)))
        new_image_path = f"resized_{os.path.basename(image_path)}"
        img.save(new_image_path)
        return new_image_path

    def _encode_image(self, image_path: str) -> str:
        if image_path.startswith("http"):
            user_agent = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/119.0.0.0 Safari/537.36 Edg/119.0.0.0"
            request_kwargs = {
                "headers": {"User-Agent": user_agent},
                "stream": True,
            }

            response = requests.get(image_path, **request_kwargs)
            response.raise_for_status()
            content_type = response.headers.get("content-type", "")

            extension = mimetypes.guess_extension(content_type)
            if extension is None:
                extension = ".download"

            fname = str(uuid.uuid4()) + extension
            download_path = os.path.abspath(os.path.join("downloads", fname))

            with open(download_path, "wb") as fh:
                for chunk in response.iter_content(chunk_size=512):
                    fh.write(chunk)

            image_path = download_path

        with open(image_path, "rb") as image_file:
            return base64.b64encode(image_file.read()).decode("utf-8")

    def forward(self, file_path: str, question: str | None = None) -> str:
        self._validate_file_type(file_path)

        start_time = time.perf_counter()
        total_tokens = 0

        if not question:
            question = "Please write a detailed caption for this image."
        try:
            mime_type, _ = mimetypes.guess_type(file_path)
            base64_image = self._encode_image(file_path)
            payload = {
                "model": "gpt-4o",
                "messages": [
                    {
                        "role": "user",
                        "content": [
                            {"type": "text", "text": question},
                            {
                                "type": "image_url",
                                "image_url": {"url": f"data:{mime_type};base64,{base64_image}"},
                            },
                        ],
                    }
                ],
                "max_tokens": 2000,
                "top_p": 0.1,
            }

            headers = {
                "Content-Type": "application/json",
                "Authorization": f"Bearer {self.gpt_key}",
            }

            response = requests.post(
                f"{self.gpt_url}/chat/completions", headers=headers, json=payload
            )
            response.raise_for_status()
            resp_json = response.json()
            # Log token usage if available
            try:
                usage = resp_json.get("usage", {})
                prompt_tokens = usage.get("prompt_tokens", 0) or 0
                completion_tokens = usage.get("completion_tokens", 0) or 0
                total_tokens = usage.get("total_tokens", 0) or 0
                logger.info(
                    "[Tool:visual_inspector] input_tokens=%s output_tokens=%s total_tokens=%s",
                    str(prompt_tokens),
                    str(completion_tokens),
                    str(total_tokens),
                )
                # Also log in a compact format like search_tools (simple tokens display)
                logger.info(
                    "[Tool:visual_inspector] model=gpt-4o total_tokens=%s prompt_tokens=%s completion_tokens=%s",
                    str(total_tokens),
                    str(prompt_tokens),
                    str(completion_tokens),
                )
                # accumulate these tokens onto the shared model instance if available
                try:
                    if hasattr(self.model, "record_usage"):
                        self.model.record_usage(prompt_tokens, completion_tokens, total_tokens)
                except Exception:
                    pass
            except Exception:
                pass
            description = resp_json["choices"][0]["message"]["content"]

            duration_sec = time.perf_counter() - start_time
            logger.info(
                f"[Tool:visual_inspector] duration={duration_sec:.3f}s, tokens={total_tokens}"
            )

        except Exception as gpt_error:
            duration_sec = time.perf_counter() - start_time
            logger.warning(f"[Tool:visual_inspector] failed after {duration_sec:.3f}s: {gpt_error}")
            return f"Visual processing failed: {str(gpt_error)}"

        if not question.startswith("Please write a detailed caption"):
            return description
        return f"You did not provide a particular question, so here is a detailed description of the image: {description}"


class TextInspectorTool(Tool):
    name = "inspect_file_as_text"
    description = """
You cannot load files yourself: instead call this tool to read a file as markdown text and ask questions about it.
This tool handles the following file extensions: [".html", ".pdb", ".xlsx", ".xls", ".pdf", ".docx", ".ppt", ".pptx"], and all other types of text files. IT DOES NOT HANDLE IMAGES."""

    inputs = {
        "file_path": {
            "description": "The path to the file you want to read as text. Must be a '.something' file, like '.pdf'. If it is an image, use the visualizer tool instead! If it is an audio, use the audio tool instead! DO NOT use this tool for an HTML webpage: use the web_search tool instead!",
            "type": "string",
        },
        "question": {
            "description": "[Optional]: Your question, as a natural language sentence. Provide as much context as possible. Do not pass this parameter if you just want to directly return the content of the file.",
            "type": "string",
            "nullable": True,
        },
    }
    output_type = "string"
    md_converter = MarkdownConverter()

    def __init__(self, model: Model, text_limit: int):
        super().__init__()
        self.model = model
        self.text_limit = text_limit

    def jsonld_to_markdown(self, data):
        markdown = ""
        if isinstance(data, dict):
            for key, value in data.items():
                markdown += f"**{key}**: {self.jsonld_to_markdown(value)}\n"
        elif isinstance(data, list):
            for item in data:
                markdown += f"- {self.jsonld_to_markdown(item)}\n"
        else:
            markdown += str(data)
        return markdown

    def parse_pdb_file(self, file_path):
        from Bio import PDB

        parser = PDB.PDBParser(QUIET=True)
        structure = parser.get_structure("protein", file_path)

        atoms = list(structure.get_atoms())
        if len(atoms) < 2:
            return "Error: PDB file contains fewer than two atoms."

        atom1, atom2 = atoms[0], atoms[1]
        distance = atom1 - atom2

        return (
            f"First atom: {atom1.get_name()} ({atom1.coord})\n"
            f"Second atom: {atom2.get_name()} ({atom2.coord})\n"
            f"Distance: {distance:.3f} Angstroms ({distance * 100:.0f} pm)"
        )

    def extract_excel_data(self, file_path, max_rows=30):
        try:
            if file_path.lower().endswith(".xls"):
                # Legacy .xls format: read with xlrd
                return self._extract_xls_data(file_path, max_rows)
            else:
                # .xlsx format: read with openpyxl
                return self._extract_xlsx_data(file_path, max_rows)
        except Exception as e:
            raise Exception(f"Error reading Excel file: {e}") from e

    def _normalize_cell_value(self, value):
        """Convert a cell value to a JSON-serializable value."""
        from datetime import date, datetime, time

        if value is None:
            return ""
        if isinstance(value, (datetime, date)):
            return value.isoformat()
        if isinstance(value, time):
            return value.strftime("%H:%M:%S")
        return value

    def _extract_xls_data(self, file_path, max_rows=30):
        """Read a legacy .xls file with xlrd."""
        try:
            import xlrd
            from xlrd import xldate_as_datetime
        except ImportError as e:
            raise ImportError(
                "xlrd is required for reading .xls files. Install with: pip install xlrd"
            ) from e

        workbook = xlrd.open_workbook(file_path, formatting_info=True)
        result = []

        for sheet_idx in range(workbook.nsheets):
            sheet = workbook.sheet_by_index(sheet_idx)

            for row_idx in range(sheet.nrows):
                row_data = []
                for col_idx in range(sheet.ncols):
                    cell = sheet.cell(row_idx, col_idx)
                    # xlrd date cells (type 3) need explicit conversion
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        try:
                            cell_value = xldate_as_datetime(
                                cell.value, workbook.datemode
                            ).isoformat()
                        except Exception:
                            cell_value = str(cell.value)
                    else:
                        cell_value = self._normalize_cell_value(cell.value)

                    # Best-effort background color lookup
                    cell_color = "FFFFFF"
                    try:
                        xf_index = sheet.cell_xf_index(row_idx, col_idx)
                        xf = workbook.xf_list[xf_index]
                        bg_color_index = xf.background.pattern_colour_index
                        if bg_color_index and bg_color_index < len(workbook.colour_map):
                            rgb = workbook.colour_map.get(bg_color_index)
                            if rgb:
                                cell_color = "{:02X}{:02X}{:02X}".format(*rgb)
                    except Exception:
                        pass

                    row_data.append({"value": cell_value, "color": cell_color})
                result.append(row_data)

        return self._format_excel_result(result, max_rows)

    def _extract_xlsx_data(self, file_path, max_rows=30):
        """Read an .xlsx file with openpyxl."""
        workbook = load_workbook(file_path)
        result = []

        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]

            for row in sheet.iter_rows():
                row_data = []
                for cell in row:
                    cell_value = self._normalize_cell_value(cell.value)

                    fill = cell.fill
                    if hasattr(fill, "fgColor") and fill.fgColor.type == "rgb":
                        cell_color = fill.fgColor.rgb
                        if len(cell_color) == 8:
                            cell_color = cell_color[2:]
                    else:
                        cell_color = "FFFFFF"
                    row_data.append({"value": cell_value, "color": cell_color})
                result.append(row_data)

        return self._format_excel_result(result, max_rows)

    def _format_excel_result(self, result, max_rows=30):
        """Format extracted Excel cells as text."""
        try:
            json_data = json.dumps(result, ensure_ascii=False, indent=4)
            data_with_colors = json.loads(json_data)
            text = []
            num_rows = len(data_with_colors)
            num_cols = len(data_with_colors[0]) if data_with_colors else 0
            text.append(
                f"This is a {num_rows} rows and {num_cols} columns table. The content is shown below:"
            )
            for row in data_with_colors:
                row_text = ""
                for cell in row:
                    value = cell["value"] if cell["value"] != "" else "None"
                    color = cell["color"]
                    if color == "FFFFFF" or color == "000000":
                        row_text += f"{value} "
                    else:
                        row_text += f"{value}({color}) "
                text.append(row_text)
            return "\n".join(text)
        except Exception as e:
            return f"Error: {str(e)}"

    def forward(self, file_path, question: str | None = None) -> str:
        start_time = time.perf_counter()
        total_tokens = 0
        llm_called = False

        try:
            if ".jsonld" in file_path:
                try:
                    with open(file_path) as f:
                        data = json.load(f)
                    result_text = self.jsonld_to_markdown(data)
                    result = type(
                        "Result", (object,), {"title": file_path, "text_content": result_text}
                    )
                except Exception as e:
                    raise Exception(f"Error parsing JSON-LD file: {e!s}") from e
            else:
                result = self.md_converter.convert(file_path)

            if file_path[-4:] in [".png", ".jpg"]:
                raise Exception(
                    "Cannot use inspect_file_as_text tool with images: use visualizer instead!"
                )

            if ".zip" in file_path:
                return result.text_content

            if file_path.endswith((".ppt", ".pptx")):
                content = ""
                try:
                    ppt = Presentation(file_path)
                    for slide_number, slide in enumerate(ppt.slides, start=1):
                        content += f"=== Slide {slide_number} ===\n"
                        slide_texts = []
                        for shape in slide.shapes:
                            if hasattr(shape, "text") and shape.text.strip():
                                slide_texts.append(shape.text.strip())
                        content += "\n".join(slide_texts) + "\n\n"
                    result.text_content = content.strip()
                except Exception as e:
                    return f"Error parsing PPT file: {e}"

            if file_path.endswith((".xls", ".xlsx")):
                try:
                    full_content = self.extract_excel_data(file_path)
                    if not question:
                        return full_content
                    else:
                        result.text_content = full_content
                except Exception as e:
                    raise Exception(f"Error processing Excel file: {e}") from e

            if file_path.endswith(".pdb"):
                try:
                    pdb_info = self.parse_pdb_file(file_path)
                    if not question:
                        return f"Extracted PDB Data:\n{pdb_info}"

                    else:
                        result.text_content = pdb_info
                except Exception as e:
                    raise Exception(f"Error parsing PDB file: {e}") from e

            if not question:
                return result.text_content

            messages = [
                {
                    "role": MessageRole.SYSTEM,
                    "content": [
                        {
                            "type": "text",
                            "text": "You will have to write a short caption for this file, then answer this question:"
                            + question,
                        }
                    ],
                },
                {
                    "role": MessageRole.USER,
                    "content": [
                        {
                            "type": "text",
                            "text": "Here is the complete file:\n### "
                            + str(result.title)
                            + "\n\n"
                            + result.text_content[: self.text_limit],
                        }
                    ],
                },
                {
                    "role": MessageRole.USER,
                    "content": [
                        {
                            "type": "text",
                            "text": "Now answer the question below. Use these three headings: '1. Short answer', '2. Extremely detailed answer', '3. Additional Context on the document and question asked'."
                            + question,
                        }
                    ],
                },
            ]
            llm_called = True
            response = self.model(messages)

            # Token usage of the last model call, if available
            if hasattr(self.model, "last_input_token_count") and hasattr(
                self.model, "last_output_token_count"
            ):
                total_tokens = (self.model.last_input_token_count or 0) + (
                    self.model.last_output_token_count or 0
                )
                # Log LLM usage in compact form consistent with search_tools
                logger.info(
                    "[TextInspector Usage] total_tokens=%s input_tokens=%s output_tokens=%s",
                    str(total_tokens),
                    str(self.model.last_input_token_count or 0),
                    str(self.model.last_output_token_count or 0),
                )

            return response.content

        finally:
            duration_sec = time.perf_counter() - start_time
            logger.info(
                f"[Tool:text_inspector] duration={duration_sec:.3f}s, tokens={total_tokens}, llm_called={llm_called}"
            )


class AudioInspectorTool(Tool):
    name = "inspect_file_as_audio"
    description = """
You cannot load files directly: use this tool to process audio files and answer related questions.
This tool supports the following audio formats: [".mp3", ".m4a", ".wav", ".flac", ".ogg"]. For other file types, use the appropriate inspection tool."""

    inputs = {
        "file_path": {
            "description": "The path to the file you want to read as audio. Must be a '.something' file, like '.mp3','.m4a','.wav'. If it is an text, use the text_inspector tool instead! If it is an image, use the visualizer tool instead! DO NOT use this tool for an HTML webpage: use the web_search tool instead!",
            "type": "string",
        },
        "question": {
            "description": "[Optional]: Your question about the audio content. Provide as much context as possible. Do not pass this parameter if you just want to directly return the content of the file.",
            "type": "string",
            "nullable": True,
        },
    }
    output_type = "string"

    # Class-level cache so the Whisper model is loaded only once
    _whisper_model = None
    _whisper_model_size = None

    def __init__(
        self,
        model: Model,
        text_limit: int,
        whisper_model_size: str = "base",
        whisper_device: str = "auto",
        whisper_compute_type: str = "auto",
    ):
        """
        Args:
            model: LLM used to answer questions about the transcript.
            text_limit: Maximum number of transcript characters passed to the LLM.
            whisper_model_size: faster-whisper model size (tiny, base, small, medium, large-v2, large-v3).
            whisper_device: Device ("auto", "cuda", "cpu").
            whisper_compute_type: Compute type ("auto", "float16", "int8", "int8_float16").
        """
        super().__init__()
        self.model = model
        self.text_limit = text_limit
        self.whisper_model_size = whisper_model_size
        self.whisper_device = whisper_device
        self.whisper_compute_type = whisper_compute_type

        self._ensure_whisper_model()

    def _ensure_whisper_model(self):
        """Load the faster-whisper model on first use (cached at class level)."""
        if (
            AudioInspectorTool._whisper_model is None
            or AudioInspectorTool._whisper_model_size != self.whisper_model_size
        ):
            try:
                from faster_whisper import WhisperModel

                # Resolve "auto" device and compute type
                device = self.whisper_device
                compute_type = self.whisper_compute_type

                if device == "auto":
                    import torch

                    device = "cuda" if torch.cuda.is_available() else "cpu"

                if compute_type == "auto":
                    compute_type = "float16" if device == "cuda" else "int8"

                logger.info(
                    f"Loading faster-whisper model: {self.whisper_model_size} on {device} with {compute_type}"
                )
                AudioInspectorTool._whisper_model = WhisperModel(
                    self.whisper_model_size,
                    device=device,
                    compute_type=compute_type,
                )
                AudioInspectorTool._whisper_model_size = self.whisper_model_size
                logger.info("Faster-whisper model loaded successfully")
            except ImportError as e:
                raise ImportError(
                    "faster-whisper is not installed. Install with: pip install faster-whisper"
                ) from e
            except Exception as e:
                raise RuntimeError(f"Failed to load faster-whisper model: {e}") from e

    def _validate_file_type(self, file_path: str):
        """Validate if the file type is a supported audio format"""
        supported_formats = [".mp3", ".m4a", ".wav", ".flac", ".ogg"]
        if not any(file_path.lower().endswith(ext) for ext in supported_formats):
            raise ValueError(f"Unsupported file type. Supported formats: {supported_formats}")

    def transcribe_audio(self, file_path: str) -> str:
        """Transcribe audio using faster-whisper (local)"""
        self._ensure_whisper_model()

        try:
            segments, info = AudioInspectorTool._whisper_model.transcribe(
                file_path,
                beam_size=5,
                language=None,  # Auto-detect the language
                vad_filter=True,  # Filter out silence with VAD
            )

            transcript_parts = []
            for segment in segments:
                transcript_parts.append(segment.text.strip())

            transcript = " ".join(transcript_parts)
            # Attempt to extract any token-like info from the whisper `info` object
            token_count = None
            try:
                token_count = getattr(info, "num_tokens", None)
            except Exception:
                token_count = None
            if token_count is None and isinstance(info, dict):
                token_count = info.get("num_tokens") or info.get("tokens")

            # Fallback estimated token count from transcript length (approx. 1 token per word)
            estimated_tokens = len(transcript.split())

            if token_count is not None:
                logger.info(
                    "[Whisper Usage] duration=%s estimated_transcript_tokens=%s (reported=%s)",
                    f"{getattr(info, 'duration', None):.1f}s"
                    if hasattr(info, "duration")
                    else str(info.get("duration", "N/A")),
                    str(estimated_tokens),
                    str(token_count),
                )
            else:
                logger.info(
                    "[Whisper Usage] duration=%s estimated_transcript_tokens=%s",
                    f"{getattr(info, 'duration', None):.1f}s"
                    if hasattr(info, "duration")
                    else str(info.get("duration", "N/A")),
                    str(estimated_tokens),
                )

            return transcript
        except Exception as e:
            raise RuntimeError(f"Speech recognition failed: {str(e)}") from e

    def forward(self, file_path: str, question: str | None = None) -> str:
        self._validate_file_type(file_path)

        start_time = time.perf_counter()
        total_tokens = 0
        llm_called = False

        if not os.path.exists(file_path):
            return f"Audio processing error: File not found: {file_path}"

        try:
            transcript = self.transcribe_audio(file_path)
        except Exception as e:
            duration_sec = time.perf_counter() - start_time
            logger.warning(
                f"[Tool:audio_inspector] transcription failed after {duration_sec:.3f}s: {e}"
            )
            return f"Audio processing error: {str(e)}"

        if not transcript.strip():
            duration_sec = time.perf_counter() - start_time
            logger.info(f"[Tool:audio_inspector] no speech detected, duration={duration_sec:.3f}s")
            return "Audio transcription: (No speech detected in the audio)"

        if not question:
            duration_sec = time.perf_counter() - start_time
            logger.info(f"[Tool:audio_inspector] transcription only, duration={duration_sec:.3f}s")
            return f"Audio transcription:\n{transcript[: self.text_limit]}"

        messages = [
            {
                "role": MessageRole.SYSTEM,
                "content": [
                    {
                        "type": "text",
                        "text": f"Audio transcription:\n{transcript[: self.text_limit]}",
                    }
                ],
            },
            {
                "role": MessageRole.USER,
                "content": [
                    {
                        "type": "text",
                        "text": "Answer the following question based on the audio content using the format:\n"
                        "1. Brief answer\n2. Detailed analysis\n3. Relevant context\nQuestion: {question}",
                    }
                ],
            },
        ]

        llm_called = True
        response = self.model(messages)

        # Token usage of the last model call, if available
        if hasattr(self.model, "last_input_token_count") and hasattr(
            self.model, "last_output_token_count"
        ):
            total_tokens = (self.model.last_input_token_count or 0) + (
                self.model.last_output_token_count or 0
            )
            logger.info(
                "[AudioInspector Usage] total_tokens=%s input_tokens=%s output_tokens=%s",
                str(total_tokens),
                str(self.model.last_input_token_count or 0),
                str(self.model.last_output_token_count or 0),
            )

        duration_sec = time.perf_counter() - start_time
        logger.info(
            f"[Tool:audio_inspector] duration={duration_sec:.3f}s, tokens={total_tokens}, llm_called={llm_called}"
        )

        return response.content


__all__ = ["VisualInspectorTool", "TextInspectorTool", "AudioInspectorTool"]
